"""D3 (DuRetrieval) shared pieces: the frame, human-label batches read from annotated sheets, and the multi-batch
design-based estimator used by the D3 analysis and the batch planner.

Symbols (witness.core): frame items d (in some system's top-10), disc [S, N], g [P, N], G = [g; disc] [R, N].
Batch t: Poisson draw with known inclusion probabilities pi_t over the items not labelled in earlier batches, shown to
annotators in random order; if the first m rows of the sheet are labelled, an item drawn in batch t is labelled with
prob pi_t * m_unique / n_unique (thinning).  Gains: u0 = 1[y=2], u_half = (u0+u1)/2, u1 = 1[y>=1].
Estimator ('final' mode of scripts/eval/evaluate.py): one GREG fit H on every label, each batch contributes its HT
correction over its own frame, batches combined with weights 1/sqrt(predicted variance) fixed before each draw.
"""
import json
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd

from witness import core as W


def load_frame(cfg):
    out = cfg["out_dir"]
    frame = pd.read_parquet(f"{out}/frame.parquet")
    z = np.load(f"{out}/frame_disc.npz")
    assert (z["qi"] == frame.qi.to_numpy()).all() and (z["di"] == frame.di.to_numpy()).all()
    disc, systems, n_q = z["disc"], [str(s) for s in z["systems"]], int(z["n_q"])
    pairs, g = W.pair_coefs(disc)
    J = np.stack([frame[j].to_numpy() for j in cfg["judges"]])  # [n_judges, N]
    o = frame.in_qrels.to_numpy()
    F0, F1 = (J == 2).mean(0), (J >= 1).mean(0)
    F = [np.where(o, 1.0, F0), np.where(o, 1.0, 0.5 * F0 + 0.5 * F1), np.where(o, 1.0, F1)]  # judge plug-in
    return SimpleNamespace(cfg=cfg, frame=frame, disc=disc, systems=systems, n_q=n_q, pairs=pairs, g=g,
                           G=np.vstack([g, disc]), P=len(pairs), N=len(frame), o=o, J=J, F=F,
                           state=((J + 1) * 4 ** np.arange(len(J))[:, None]).sum(0),
                           qi=frame.qi.to_numpy(),
                           pos={(int(a), int(b)): i for i, (a, b) in enumerate(zip(frame.qi, frame.di))})


def batch_pi(D, stage):
    """Design inclusion probabilities of batch `stage` over the frame (0 outside its frame)."""
    if stage == 1:
        return D.frame.pi.to_numpy().astype(float)
    return np.load(f"{D.cfg['out_dir']}/batch{stage}_pi.npy")


def sheet_labels(path):
    """Sheet row -> human grade of one batch: the annotated xlsx (sheet 标注, columns 序号 / 标签, blank rows skipped)
    or the released label CSV of that batch (columns row, label; outputs/d3/labels/ of the data-artifact bundle)."""
    if path.endswith(".csv"):
        df = pd.read_csv(path)
        labels = {int(r): int(v) for r, v in zip(df["row"], df["label"])}
        assert len(labels) == len(df)
    else:
        ws = openpyxl.load_workbook(path, data_only=True)["标注"]
        head = [c.value for c in ws[1]]
        c_row, c_lab = head.index("序号"), head.index("标签")
        labels = {int(r[c_row]): int(str(r[c_lab]).strip()) for r in ws.iter_rows(min_row=2, values_only=True)
                  if r[c_lab] is not None and str(r[c_lab]).strip() != ""}
    assert labels and all(v in (0, 1, 2) for v in labels.values())
    return labels


def read_batch(D, stage, sheet):
    """Labelled items of a batch.  Returns SimpleNamespace(stage, items [k], y [k], pi [N] thinned, agree [r])."""
    key = json.load(open(f"{D.cfg['out_dir']}/batch{stage}_key.json", encoding="utf8"))
    labels = sheet_labels(sheet)
    m = max(labels)
    assert all(k in labels for k in range(1, m + 1)), "labelled rows must be a prefix of the sheet"
    rows = {k["row"]: k for k in key["rows"]}
    first, agree = {}, []
    for k in range(1, m + 1):
        kr = rows[k]
        it = D.pos[(kr["qi"], kr["di"])]
        if kr["repeat"]:
            if it in first:
                agree.append(first[it] == labels[k])
            continue
        first[it] = labels[k]
    pi = batch_pi(D, stage) * (len(first) / key["n_unique"])
    items = np.array(sorted(first))
    assert (pi[items] > 0).all()
    return SimpleNamespace(stage=stage, items=items, y=np.array([first[i] for i in items]), pi=pi, agree=agree,
                           m_rows=m, n_unique=key["n_unique"])


def gains(y):
    return [(y == 2).astype(float), 0.5 * (y == 2) + 0.5 * (y >= 1), (y >= 1).astype(float)]


def design_columns(D, b, extras, f):
    """GREG columns: qrels stratum, U x judge joint state, the base predictor f [N] of block b on U (the judge
    plug-in D.F[b] in the main analysis), optional auxiliaries (centred on U)."""
    o, U = D.o, ~D.o
    cols = [o.astype(float)] + [(U & (D.state == lev)).astype(float) for lev in np.unique(D.state)]
    cols.append(np.where(o, 0.0, f))
    b0 = [0.0] * (len(cols) - 1) + [1.0]
    sysm = np.array([s.startswith("rerank") for s in D.systems])
    lexm = np.array([s.startswith("bm25") for s in D.systems])
    for x in extras:
        if x == "fam":
            new = [(D.disc[lexm] > 0).mean(0), (D.disc[~lexm & ~sysm] > 0).mean(0), (D.disc[sysm] > 0).mean(0)]
        else:
            assert x == "sys"
            new = [c for s in range(len(D.systems)) for c in ((D.disc[s] > 0).astype(float), D.disc[s])]
        for c in new:
            cols.append(np.where(U, c - c[U].mean(), 0.0))
            b0.append(0.0)
    return np.column_stack(cols), np.array(b0)


def write_sheet(D, rows, path, repeat_share, rng):
    """Blinded annotation sheet in random order with exact repeats in the second half.  rows: frame indices.
    Returns the key list (row, qi, di, stratum, repeat) for batch<stage>_key.json."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as ILLEGAL
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = np.array(rows)
    rng.shuffle(rows)
    order = list(rows)
    n_rep = int(round(repeat_share * len(rows)))
    for r in rng.choice(rows[: len(rows) // 2], size=n_rep, replace=False):
        order.insert(int(rng.integers(len(order) // 2, len(order) + 1)), r)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "说明"
    for line in (
        "任务：判断每个“段落”是否包含能直接回答“查询”所表达的信息需求的内容。",
        "0 = 不相关，或只是主题相关但不含答案。",
        "1 = 部分回答：含相关信息，但不完整、不够具体或需要推断。",
        "2 = 直接且完整地回答。",
        "只依据给定段落本身作答，不要使用搜索引擎或外部知识。可以用常识理解查询所指的对象（如缩写、歌名），但答案必须在段落里。",
        "在“标签”列选择 0、1 或 2。拿不准时选第一直觉，并在“备注”写一两个字（如“拿不准”“查询有歧义”“段落截断”）。",
        "少数条目会重复出现，用来检查一致性，请照常独立判断，不要回头翻之前的选择。",
        "行序是随机的：从第 1 行起连续标到任意一行都可用，不必标完；请不要跳行。",
        "多人分工时按行号切成连续的段，每人一段，并在“标注人”列写自己的代号。",
        "完成后保持 xlsx 格式保存，不要改动“序号”列。",
    ):
        ws0.append([line])
    ws0.column_dimensions["A"].width = 120
    ws = wb.create_sheet("标注")
    ws.append(["序号", "查询", "段落", "标签", "备注", "标注人"])
    for c in "ABCDEF":
        ws[f"{c}1"].font = Font(bold=True)
    key, seen = [], set()
    for n, r in enumerate(order, start=1):
        ws.append([n, ILLEGAL.sub("", D.frame["query"].iloc[r]), ILLEGAL.sub("", D.frame["passage"].iloc[r]), None, None, None])
        key.append(dict(row=n, qi=int(D.frame.qi.iloc[r]), di=int(D.frame.di.iloc[r]),
                        stratum="L" if D.o[r] else "U", repeat=bool(r in seen)))
        seen.add(r)
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(order) + 1}")
    for c, wdt in zip("ABCDEF", (6, 28, 110, 8, 16, 10)):
        ws.column_dimensions[c].width = wdt
    for row in ws.iter_rows(min_row=2):
        for cell in row[1:3]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return key, len(rows)


def working_var(D, b, batches):
    """Per-item working Var(u_b | x) for planning and stage weights: before any label the judge-spread model
    f(1-f) + .05 (the batch-1 design); with labels, the 1/pi-weighted mean of (u - F)^2 per (stratum x judge state)
    cell, cells with < 5 labels pooled within their stratum."""
    if not batches:
        f = D.F[b]
        return f * (1 - f) + 0.05
    items = np.concatenate([bt.items for bt in batches])
    u = np.concatenate([gains(bt.y)[b] for bt in batches])
    w = np.concatenate([1.0 / bt.pi[bt.items] for bt in batches])
    r2 = (u - D.F[b][items]) ** 2
    cell = D.state * 2 + D.o.astype(int)
    s2 = np.zeros(D.N)
    for strat in (False, True):
        ms = D.o[items] == strat
        pooled = np.average(r2[ms], weights=w[ms])
        for c in np.unique(cell[D.o == strat]):
            mc = ms & (cell[items] == c)
            s2[cell == c] = np.average(r2[mc], weights=w[mc]) if mc.sum() >= 5 else pooled
    return np.maximum(s2, 1e-4)


def design_estimate(D, batches, extras, ridge, F):
    """Multi-batch design-based estimator ('final' GREG fit, stage weights 1/sqrt(predicted var) fixed before each draw).
    F: base predictors per block (rho = 0, 1/2, 1), each [N] (D.F, the judge plug-in, in the main analysis), or a list
    with one such triple per batch (adaptive witness: batch t's predictor may depend on the labels of batches < t only,
    never on its own).  The stage weights use the judge-based working variance whatever F is, so they stay fixed
    before each draw.  Returns est, var: lists over the blocks (rho = 0, 1/2, 1), each [R]."""
    items = np.concatenate([bt.items for bt in batches])
    pi_lab = np.concatenate([bt.pi[bt.items] for bt in batches])
    per_batch = isinstance(F[0], (list, tuple))
    Fb = F if per_batch else [F] * len(batches)
    assert len(Fb) == len(batches)
    est_out, var_out = [], []
    for b in range(3):
        u_all = np.zeros(D.N)
        u_all[items] = np.concatenate([gains(bt.y)[b] for bt in batches])
        Hs = []
        for t in range(len(batches)):
            Zc, b0 = design_columns(D, b, extras, Fb[t][b])
            Hs.append(W.greg_fit(D.G, Zc, u_all, items, pi_lab, b0, ridge) @ Zc.T)  # [R, N]
        num = den = vnum = 0.0
        known = np.zeros(D.N, bool)
        for t, bt in enumerate(batches):
            xi = np.zeros(D.N, bool)
            xi[bt.items] = True
            assert not (known & (bt.pi > 0)).any() and not (xi & known).any()
            e_t, v_t = W.stage_estimate(D.G, known * u_all, np.zeros(D.N, bool), Hs[t], u_all, xi, bt.pi, D.n_q)
            m2 = working_var(D, b, batches[:t])[None, :] + 0.0 * Hs[t]  # [R, N]
            om = 1.0 / np.sqrt(np.maximum(W.predicted_var(D.G, bt.pi, m2, D.n_q), 1e-16))
            num, den, vnum = num + om * e_t, den + om, vnum + om ** 2 * v_t
            known |= xi
        est_out.append(num / den)
        var_out.append(vnum / den ** 2)
    return est_out, var_out
